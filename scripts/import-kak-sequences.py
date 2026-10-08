"""Extract the supplied DKPP sub-activity sequence workbook without editing it."""
import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("src/data/kak-sub-kegiatan-2027.json"))
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    try:
        sheet = workbook["Sub Kegiatan"]
        if list(next(sheet.iter_rows(min_row=5, max_row=5, values_only=True))) != ["No", "Kode Sub Kegiatan", "Nama Sub Kegiatan"]:
            raise ValueError("Kolom daftar sub kegiatan berubah; periksa sumber sebelum impor.")
        if "Dinas Ketahanan Pangan dan Pertanian" not in str(sheet["A2"].value):
            raise ValueError("Daftar bukan untuk DKPP.")
        year = re.search(r"Tahun Anggaran (\d{4})", str(sheet["A3"].value))
        if not year:
            raise ValueError("Tahun anggaran sumber tidak ditemukan.")
        entries = []
        for row, values in enumerate(sheet.iter_rows(min_row=6, max_col=3, values_only=True), 6):
            if all(value is None for value in values):
                continue
            number, code, name = values
            if type(number) is not int or number <= 0 or not isinstance(code, str) or not re.fullmatch(r"\d\.\d{2}\.\d{2}\.\d\.\d{2}\.\d{4}", code) or not isinstance(name, str) or not name.strip():
                raise ValueError(f"Baris {row} tidak valid.")
            entries.append({"number": number, "code": code, "name": name, "row": row})
        if not entries or len({entry["number"] for entry in entries}) != len(entries) or len({entry["code"] for entry in entries}) != len(entries):
            raise ValueError("Daftar kosong atau memuat nomor/kode ganda.")
        if [entry["number"] for entry in entries] != list(range(1, len(entries) + 1)):
            raise ValueError("Urutan nomor sumber perlu diperiksa.")
        data = {
            "id": f"kak-sub-kegiatan-dkpp-{year[1]}",
            "title": f"Daftar nomor urut sub kegiatan DKPP TA {year[1]}",
            "year": int(year[1]),
            "fileName": args.workbook.name,
            "sheet": sheet.title,
            "source": str(sheet["A3"].value),
            "sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
            "entries": entries,
        }
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"Imported {len(entries)} sub activities for TA {year[1]} into {args.output}")
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
